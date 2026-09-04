from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "BACKEND_DEPLOYMENT.md"
OUTPUT = ROOT / "Инструкция_развертывания_Вектор.docx"


def set_cell_shading(paragraph, color):
    props = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    props.append(shading)


def set_run_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text_with_inline_code(doc, text, style=None):
    paragraph = doc.add_paragraph(style=style)
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(part[1:-1] if part.startswith("`") else part)
        if part.startswith("`"):
            set_run_font(run, "Consolas", 9, color=(31, 78, 121))
    return paragraph


def add_code_block(doc, lines):
    for line in lines:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Cm(0.45)
        paragraph.paragraph_format.right_indent = Cm(0.25)
        set_cell_shading(paragraph, "F2F5F8")
        run = paragraph.add_run(line)
        set_run_font(run, "Consolas", 8.2, color=(31, 45, 61))


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = document.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(22)
    title.font.color.rgb = RGBColor(0, 0, 0)

    for style_name, size in (("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)):
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)

    if "Code Block" not in document.styles:
        code = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font.size = Pt(8.2)


def build_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    configure_styles(document)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines = []
    first_title = True

    for raw_line in lines:
        line = raw_line.rstrip()
        if line.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line or line == "---":
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if first_title and level == 1:
                paragraph = document.add_paragraph(text, style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                first_title = False
                add_text_with_inline_code(
                    document,
                    "Практическая инструкция по запуску серверной версии, настройке HTTPS, резервному копированию и обновлению приложения.",
                )
            else:
                document.add_paragraph(text, style=f"Heading {level}")
            continue

        if line.startswith("- "):
            add_text_with_inline_code(document, line[2:], style="List Bullet")
            continue
        if re.match(r"^\d+\. ", line):
            add_text_with_inline_code(document, re.sub(r"^\d+\. ", "", line), style="List Number")
            continue

        add_text_with_inline_code(document, line)

    if in_code:
        raise RuntimeError("Unclosed code block in source Markdown")
    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
