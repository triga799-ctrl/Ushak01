from pathlib import Path

from docx import Document


document_path = Path(__file__).resolve().parents[1] / "Инструкция_развертывания_Вектор.docx"
document = Document(document_path)
content = "\n".join(paragraph.text for paragraph in document.paragraphs)

assert document_path.stat().st_size > 10_000
assert "Как посмотреть содержимое файла backup" in content
assert "Первый запуск" in content
assert "Построчное пояснение команд развертывания" in content

print(f"file_bytes={document_path.stat().st_size}")
print(f"paragraphs={len(document.paragraphs)}")
print("content_check=ok")
